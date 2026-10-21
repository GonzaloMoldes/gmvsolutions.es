# scripts/descargables/generar-plantilla-polivalencia.py
# Genera la plantilla Excel de matriz de polivalencia en src/descargables/.
# Se regenera desde aqui en vez de editarse a mano.
#
#   python3 scripts/descargables/generar-plantilla-polivalencia.py
#   python3 <skill xlsx>/scripts/recalc.py src/descargables/plantilla-matriz-polivalencia-reelevo.xlsx
#
# Diseno: el usuario rellena una sola hoja (Matriz: operarios x puestos, nivel 0-4).
# Los totales por puesto y por operario y el Resumen de riesgo se calculan solos.
# Riesgo = puesto con un solo operario autonomo (nivel >= 3); sin cobertura = ninguno.
# Mismos niveles y misma regla que el generador de /recursos/matriz-de-polivalencia/.
#
# Datos de ejemplo ficticios (nombres y puestos inventados). Sin datos de clientes.
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / 'src' / 'descargables' / 'plantilla-matriz-polivalencia-reelevo.xlsx'

# Capacidad de la plantilla.
MAX_OPER = 60
MAX_PUESTOS = 30

NARANJA = 'F4521E'
TINTA = '0A0A0A'
GRIS = '6B6B6B'
GRIS_FONDO = 'F2F2F2'
AMARILLO = 'FFF4CC'   # celdas a rellenar
F = 'Arial'

NIVELES = [
    (0, 'No formado', 'No ha hecho el puesto ni ha recibido formación.', 'FFFFFF', TINTA),
    (1, 'En formación', 'Está aprendiendo. No puede quedarse solo en el puesto.', 'FDE2D8', '8A2A0A'),
    (2, 'Con supervisión', 'Lo hace, pero necesita a alguien cerca para dudas o cambios.', 'FFF0C2', '6B4E00'),
    (3, 'Autónomo', 'Cubre el puesto solo, con calidad y seguridad, en cualquier turno.', 'D9F2E1', '1E6B3A'),
    (4, 'Puede formar', 'Autónomo y capaz de enseñar el puesto a otra persona.', '9FDDB4', '0F4A25'),
]

def font(size=10, bold=False, color=TINTA, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)

borde = Border(*(Side(style='thin', color='D9D9D9'),) * 4)
cab_fill = PatternFill('solid', fgColor=NARANJA)
gris_fill = PatternFill('solid', fgColor=GRIS_FONDO)
input_fill = PatternFill('solid', fgColor=AMARILLO)
wrap = Alignment(wrap_text=True, vertical='top')
centro = Alignment(horizontal='center', vertical='center', wrap_text=True)

def cab(c, texto):
    c.value = texto
    c.font = font(10, True, 'FFFFFF'); c.fill = cab_fill; c.alignment = centro; c.border = borde

wb = Workbook()

# ── Como usar ────────────────────────────────────────────────────────────────
ws = wb.active; ws.title = 'Cómo usar'
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 110
ws['B1'] = 'Plantilla de matriz de polivalencia'; ws['B1'].font = font(16, True, NARANJA)
ws['B2'] = 'REELEVO · gmvsolutions.es · Datos de ejemplo ficticios: bórralos y pon los tuyos.'; ws['B2'].font = font(9, False, GRIS, True)
texto = [
    ('p', 'Una matriz de polivalencia cruza a los operarios con los puestos de la planta y dice, para cada combinación, con qué nivel de autonomía puede trabajar esa persona en ese puesto. Sirve para cubrir bajas y vacaciones, planificar turnos y decidir a quién formar primero.'),
    ('h', 'Cómo se rellena'),
    ('p', '1. En la hoja Matriz, escribe los puestos en la fila de cabecera (celdas amarillas, hasta 30) y los operarios en la primera columna (hasta 60), con su turno.'),
    ('p', '2. En cada cruce, pon el nivel de 0 a 4. Si lo dejas vacío cuenta como 0. Los colores salen solos.'),
    ('p', '3. Lo demás se calcula: debajo de cada puesto, cuántos autónomos tiene y su estado; a la derecha de cada operario, en cuántos puestos es autónomo; y en Resumen, qué puestos están en riesgo.'),
    ('h', 'Los cinco niveles'),
    *[('p', f'{n} · {nombre}: {desc}') for n, nombre, desc, _, _ in NIVELES],
    ('h', 'Qué es un puesto en riesgo'),
    ('p', 'Un puesto con un solo operario autónomo (nivel 3 o 4). Si esa persona falta, el puesto se para o lo cubre alguien que no está preparado. Sin cobertura = ningún autónomo. La regla práctica de muchas plantas es tener al menos tres personas autónomas por puesto crítico, repartidas entre turnos.'),
    ('h', 'Cuándo se queda corta una hoja de cálculo'),
    ('p', 'Construirla es fácil; mantenerla, no. A los pocos meses deja de reflejar la realidad: nadie la actualiza cuando alguien aprende un puesto nuevo o lleva tiempo sin hacerlo, y el nivel 3 de hace un año ya no es el de hoy. En REELEVO la matriz se alimenta de cada ejecución registrada en el puesto y avisa de los puestos que se quedan sin cobertura.'),
    ('p', 'Más información: gmvsolutions.es/gestion-competencias/'),
]
fila = 4
for tipo, t in texto:
    c = ws.cell(row=fila, column=2, value=t)
    if tipo == 'h':
        c.font = font(12, True); fila += 1
    else:
        c.font = font(10); c.alignment = wrap
        ws.row_dimensions[fila].height = max(16, 15 * (len(t) // 120 + 1))
    fila += 1
ws.sheet_view.showGridLines = False

# ── Matriz ──────────────────────────────────────────────────────────────────
wm = wb.create_sheet('Matriz')
wm['A1'] = 'Matriz de polivalencia'; wm['A1'].font = font(14, True, NARANJA)
wm['A2'] = 'Rellena lo amarillo: puestos, operarios, turno y niveles (0 No formado · 1 En formación · 2 Con supervisión · 3 Autónomo · 4 Puede formar).'
wm['A2'].font = font(9, False, GRIS, True)

H = 4                        # fila de cabecera
R1, R2 = H + 1, H + MAX_OPER  # filas de operarios
C1 = 3                       # primera columna de puestos (C)
C2 = C1 + MAX_PUESTOS - 1
CA, CF = C2 + 1, C2 + 2      # columnas de totales por operario
T_AUT, T_FOR, T_EST = R2 + 2, R2 + 3, R2 + 4  # filas de totales por puesto

cab(wm.cell(row=H, column=1), 'Operario')
cab(wm.cell(row=H, column=2), 'Turno')
wm.column_dimensions['A'].width = 24
wm.column_dimensions['B'].width = 10
for c in range(C1, C2 + 1):
    cell = wm.cell(row=H, column=c)
    cell.fill = input_fill; cell.font = font(9, True); cell.alignment = centro; cell.border = borde
    wm.column_dimensions[get_column_letter(c)].width = 12
cab(wm.cell(row=H, column=CA), 'Puestos en los que es autónomo')
cab(wm.cell(row=H, column=CF), 'Puestos que puede enseñar')
wm.column_dimensions[get_column_letter(CA)].width = 14
wm.column_dimensions[get_column_letter(CF)].width = 14
wm.row_dimensions[H].height = 48

for r in range(R1, R2 + 1):
    for c in range(1, C2 + 1):
        cell = wm.cell(row=r, column=c)
        cell.fill = input_fill; cell.border = borde
        cell.font = font(10); cell.alignment = centro if c >= C1 else Alignment(vertical='center')
    rng = f'{get_column_letter(C1)}{r}:{get_column_letter(C2)}{r}'
    for c, f in ((CA, f'=IF($A{r}="","",COUNTIF({rng},">=3"))'), (CF, f'=IF($A{r}="","",COUNTIF({rng},4))')):
        cell = wm.cell(row=r, column=c, value=f)
        cell.fill = gris_fill; cell.border = borde; cell.font = font(10, True); cell.alignment = centro

for fila_t, etiqueta in ((T_AUT, 'Autónomos (nivel 3 o 4)'), (T_FOR, 'Pueden formar (nivel 4)'), (T_EST, 'Estado del puesto')):
    c = wm.cell(row=fila_t, column=1, value=etiqueta); c.font = font(10, True); c.fill = gris_fill; c.border = borde
    wm.cell(row=fila_t, column=2).fill = gris_fill
for c in range(C1, C2 + 1):
    L = get_column_letter(c)
    col = f'{L}{R1}:{L}{R2}'
    vals = {
        T_AUT: f'=IF({L}{H}="","",COUNTIF({col},">=3"))',
        T_FOR: f'=IF({L}{H}="","",COUNTIF({col},4))',
        T_EST: f'=IF({L}{H}="","",IF({L}{T_AUT}=0,"Sin cobertura",IF({L}{T_AUT}=1,"Riesgo","OK")))',
    }
    for fila_t, f in vals.items():
        cell = wm.cell(row=fila_t, column=c, value=f)
        cell.fill = gris_fill; cell.border = borde; cell.font = font(10, True); cell.alignment = centro

dv = DataValidation(type='whole', operator='between', formula1='0', formula2='4', allow_blank=True,
                    errorTitle='Nivel no válido', error='Escribe un nivel de 0 a 4.')
dv.add(f'{get_column_letter(C1)}{R1}:{get_column_letter(C2)}{R2}')
dv_turno = DataValidation(type='list', formula1='"Mañana,Tarde,Noche,Central,Rotativo"', allow_blank=True)
dv_turno.add(f'B{R1}:B{R2}')
wm.add_data_validation(dv); wm.add_data_validation(dv_turno)

niv = f'{get_column_letter(C1)}{R1}:{get_column_letter(C2)}{R2}'
for n, _, _, bg, fg in NIVELES[1:]:
    wm.conditional_formatting.add(niv, CellIsRule(operator='equal', formula=[str(n)],
        fill=PatternFill('solid', fgColor=bg), font=Font(name=F, size=10, bold=n >= 3, color=fg)))
est = f'{get_column_letter(C1)}{T_EST}:{get_column_letter(C2)}{T_EST}'
for txt, bg, fg in (('Sin cobertura', 'F4521E', 'FFFFFF'), ('Riesgo', 'FDE2D8', 'B3360C'), ('OK', 'D9F2E1', '1E6B3A')):
    wm.conditional_formatting.add(est, CellIsRule(operator='equal', formula=[f'"{txt}"'],
        fill=PatternFill('solid', fgColor=bg), font=Font(name=F, size=10, bold=True, color=fg)))
wm.freeze_panes = wm.cell(row=R1, column=C1)

# Datos de ejemplo (ficticios)
puestos = ['Inyectora 1', 'Inyectora 2', 'Montaje', 'Control calidad', 'Envasado', 'Carretilla']
operarios = [
    ('Ana R.', 'Mañana',  [4, 3, 2, 1, 3, 0]),
    ('Luis M.', 'Mañana', [3, 2, 4, 0, 2, 3]),
    ('Carmen S.', 'Tarde', [2, 1, 3, 3, 4, 0]),
    ('Javier T.', 'Tarde', [1, 0, 3, 0, 3, 3]),
    ('Marta G.', 'Noche',  [0, 0, 2, 0, 3, 0]),
    ('Pedro L.', 'Noche',  [2, 1, 1, 0, 2, 4]),
    ('Sara P.', 'Mañana',  [1, 0, 3, 2, 3, 0]),
    ('Diego F.', 'Tarde',  [0, 0, 1, 0, 1, 0]),
]
for i, p in enumerate(puestos):
    wm.cell(row=H, column=C1 + i, value=p)
for j, (nombre, turno, niveles) in enumerate(operarios):
    wm.cell(row=R1 + j, column=1, value=nombre)
    wm.cell(row=R1 + j, column=2, value=turno)
    for i, n in enumerate(niveles):
        wm.cell(row=R1 + j, column=C1 + i, value=n)

# ── Resumen ─────────────────────────────────────────────────────────────────
wr = wb.create_sheet('Resumen', 1)
wr['A1'] = 'Resumen de cobertura (se calcula solo)'; wr['A1'].font = font(14, True, NARANJA)
wr['A2'] = 'No escribas aquí. Riesgo = un solo autónomo en el puesto. Sin cobertura = ninguno.'; wr['A2'].font = font(9, False, GRIS, True)
wr.column_dimensions['A'].width = 40
wr.column_dimensions['B'].width = 14
wr.column_dimensions['C'].width = 14
wr.column_dimensions['D'].width = 18
EST = f"Matriz!${get_column_letter(C1)}${T_EST}:${get_column_letter(C2)}${T_EST}"
kpis = [
    ('Puestos en la matriz', f'=COUNTA(Matriz!${get_column_letter(C1)}${H}:${get_column_letter(C2)}${H})'),
    ('Operarios en la matriz', f'=COUNTA(Matriz!$A${R1}:$A${R2})'),
    ('Puestos sin cobertura (ningún autónomo)', f'=COUNTIF({EST},"Sin cobertura")'),
    ('Puestos en riesgo (un solo autónomo)', f'=COUNTIF({EST},"Riesgo")'),
    ('Puestos con al menos dos autónomos', f'=COUNTIF({EST},"OK")'),
]
for k, (et, f) in enumerate(kpis):
    r = 4 + k
    a = wr.cell(row=r, column=1, value=et); a.font = font(10, True); a.border = borde
    b = wr.cell(row=r, column=2, value=f); b.font = font(12, True); b.alignment = centro; b.border = borde; b.fill = gris_fill
for r in (6, 7):
    wr.conditional_formatting.add(f'B{r}', CellIsRule(operator='greaterThan', formula=['0'],
        fill=PatternFill('solid', fgColor='FDE2D8'), font=Font(name=F, size=12, bold=True, color='B3360C')))

P0 = 11
for i, t in enumerate(['Puesto', 'Autónomos', 'Pueden formar', 'Estado']):
    cab(wr.cell(row=P0, column=1 + i), t)
for k in range(MAX_PUESTOS):
    r = P0 + 1 + k
    L = get_column_letter(C1 + k)
    for i, f in enumerate([f'=Matriz!{L}{H}&""', f'=Matriz!{L}{T_AUT}&""', f'=Matriz!{L}{T_FOR}&""', f'=Matriz!{L}{T_EST}&""']):
        c = wr.cell(row=r, column=1 + i, value=f)
        c.font = font(10, i == 3); c.border = borde
        c.alignment = Alignment(vertical='center') if i == 0 else centro
rng_est = f'D{P0 + 1}:D{P0 + MAX_PUESTOS}'
for txt, bg, fg in (('Sin cobertura', 'F4521E', 'FFFFFF'), ('Riesgo', 'FDE2D8', 'B3360C'), ('OK', 'D9F2E1', '1E6B3A')):
    wr.conditional_formatting.add(rng_est, CellIsRule(operator='equal', formula=[f'"{txt}"'],
        fill=PatternFill('solid', fgColor=bg), font=Font(name=F, size=10, bold=True, color=fg)))

wb.active = 0
wb.save(OUT)
print('Generado', OUT)
