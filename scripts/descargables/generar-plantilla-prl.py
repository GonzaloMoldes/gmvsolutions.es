# scripts/descargables/generar-plantilla-prl.py
# Genera la plantilla Excel de gestion documental de PRL e instrucciones de trabajo
# en src/descargables/. Se regenera desde aqui en vez de editarse a mano.
#
#   python3 scripts/descargables/generar-plantilla-prl.py
#   python3 <skill xlsx>/scripts/recalc.py src/descargables/plantilla-gestion-documental-prl-reelevo.xlsx
#
# Diseno: el usuario rellena tres hojas (Instrucciones, Trabajadores, Acuses) y la
# Matriz y el Resumen se calculan solos. La matriz no se teclea: es justo la tabla que
# en papel se descontrola, asi que aqui sale de cruzar el registro de acuses con la
# version vigente, el alcance y el reciclaje de cada instruccion.
#
# Datos de ejemplo ficticios (nombres y codigos inventados). Sin datos de clientes.
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / 'src' / 'descargables' / 'plantilla-gestion-documental-prl-reelevo.xlsx'

# Capacidad de la plantilla. La matriz se dimensiona a esto.
MAX_INSTR = 40
MAX_TRAB = 120
MAX_ACUSES = 3000

NARANJA = 'F4521E'
TINTA = '0A0A0A'
GRIS = '6B6B6B'
GRIS_FONDO = 'F2F2F2'
AMARILLO = 'FFF4CC'   # celdas a rellenar
F = 'Arial'

def font(size=10, bold=False, color=TINTA, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)

borde = Border(*(Side(style='thin', color='D9D9D9'),) * 4)
cab_fill = PatternFill('solid', fgColor=NARANJA)
gris_fill = PatternFill('solid', fgColor=GRIS_FONDO)
input_fill = PatternFill('solid', fgColor=AMARILLO)
wrap = Alignment(wrap_text=True, vertical='top')
centro = Alignment(horizontal='center', vertical='center', wrap_text=True)

def cabecera(ws, fila, textos, anchos):
    for i, (t, w) in enumerate(zip(textos, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font = font(10, True, 'FFFFFF'); c.fill = cab_fill; c.alignment = centro; c.border = borde
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[fila].height = 32

def titulo(ws, texto, sub):
    ws['A1'] = texto; ws['A1'].font = font(14, True, NARANJA)
    ws['A2'] = sub; ws['A2'].font = font(9, False, GRIS, True)

def zona_input(ws, f1, f2, c1, c2):
    for r in range(f1, f2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = input_fill; cell.border = borde; cell.font = font(10)

wb = Workbook()

# ── 1. Como usar ─────────────────────────────────────────────────────────────
ws = wb.active; ws.title = 'Cómo usar'
ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 110
ws['B1'] = 'Plantilla de gestión documental de PRL e instrucciones de trabajo'; ws['B1'].font = font(16, True, NARANJA)
ws['B2'] = 'REELEVO · gmvsolutions.es · Plantilla gratuita. Los datos que trae son un ejemplo ficticio: bórralos y pon los tuyos.'; ws['B2'].font = font(9, False, GRIS, True)
lineas = [
    ('h', 'Para qué sirve'),
    ('p', 'Saber, para cada trabajador y cada instrucción de PRL, de trabajo (IT), de calidad o política, si conoce la versión vigente, qué tiene pendiente y qué caduca. Es la tabla que en papel se descontrola: aquí no se rellena a mano, se calcula.'),
    ('h', 'Qué rellenas tú (celdas amarillas)'),
    ('p', '1. Instrucciones: una fila por instrucción, con su versión vigente (número entero: 1, 2, 3…), a quién afecta y cada cuántos meses hay que reciclarla (0 si no caduca).'),
    ('p', '2. Trabajadores: una fila por persona, con su puesto, turno y si es de ETT. Marca «No» en Activo cuando alguien se vaya; no borres la fila.'),
    ('p', '3. Acuses: una fila cada vez que alguien lee y acepta una versión. Es el registro: no se modifica, solo se añaden filas.'),
    ('h', 'Qué se calcula solo'),
    ('p', '· Matriz: el estado de cada persona en cada instrucción. OK, Pendiente (no ha aceptado la versión vigente), Caducado (pasó su reciclaje), Caduca pronto (en los próximos 30 días) o «—» si no le afecta.'),
    ('p', '· Resumen: combinaciones a controlar, pendientes, caducados y porcentaje al día. Los pendientes por persona salen en Trabajadores, y los pendientes por instrucción en Instrucciones.'),
    ('h', 'Cuando cambia una instrucción'),
    ('p', 'Sube el número de versión en la hoja Instrucciones. Todos los que la tenían aceptada pasan solos a «Pendiente» hasta que registres su nuevo acuse. Los acuses antiguos se quedan como histórico: así puedes demostrar qué versión conocía cada persona en cada fecha.'),
    ('h', 'Fecha de referencia'),
    ('p', 'El Resumen usa la fecha de hoy. Si quieres ver cómo estabas en otra fecha (por ejemplo, el día de una auditoría), escríbela en la celda amarilla del Resumen.'),
    ('h', 'Hasta dónde llega esta plantilla'),
    ('p', 'Funciona bien con pocas decenas de instrucciones y una plantilla estable. Se queda corta cuando: (1) los acuses se registran a mano desde hojas de firmas y alguien tiene que pasarlos; (2) hay varios turnos o personal de ETT que nunca coincide con las sesiones; (3) las revisiones son frecuentes y cada una deja a decenas de personas pendientes; (4) necesitas que nadie trabaje con una versión que no ha aceptado, no solo saberlo después.'),
    ('p', 'Por qué pasa y cómo se resuelve: gmvsolutions.es/blog/gestion-instrucciones-prl-e-it/'),
    ('h', 'Aviso'),
    ('p', 'Es una herramienta práctica de control interno, no asesoramiento jurídico. Contrasta siempre los requisitos con la Ley 31/1995, tu evaluación de riesgos y tu servicio de prevención.'),
]
r = 4
for tipo, txt in lineas:
    c = ws.cell(row=r, column=2, value=txt)
    if tipo == 'h':
        c.font = font(11, True, NARANJA); r += 1
    else:
        c.font = font(10); c.alignment = wrap
        ws.row_dimensions[r].height = 15 * max(1, len(txt) // 105 + 1); r += 1
        if lineas.index((tipo, txt)) + 1 < len(lineas) and lineas[lineas.index((tipo, txt)) + 1][0] == 'h': r += 1
ws['B' + str(r + 1)] = 'Leyenda:'; ws['B' + str(r + 1)].font = font(10, True)
ws['B' + str(r + 2)] = 'Celda amarilla = la rellenas tú.   Celda blanca o gris = fórmula, no la toques.'
ws['B' + str(r + 2)].font = font(10); ws['B' + str(r + 2)].fill = input_fill
ws.sheet_view.showGridLines = False

# ── 2. Instrucciones ─────────────────────────────────────────────────────────
wi = wb.create_sheet('Instrucciones')
titulo(wi, 'Instrucciones', 'Una fila por instrucción. Versión vigente en número entero. Alcance: «Toda la plantilla» o «Por puesto» (y escribe los puestos separados por ;).')
I0 = 4  # fila de cabecera
cabecera(wi, I0, ['Código', 'Tipo', 'Título', 'Versión vigente', 'Fecha de la versión', 'Responsable', 'Alcance', 'Puestos afectados (separados por ;)', 'Reciclaje (meses, 0 = no caduca)', 'Afectados', 'Pendientes', 'Caducados', '% al día'],
         [10, 12, 40, 10, 13, 18, 16, 34, 13, 11, 11, 11, 10])
ejemplo_instr = [
    ('PRL-01', 'PRL', 'Uso de EPI en planta', 3, date(2026, 9, 1), 'Prevención', 'Toda la plantilla', '', 12),
    ('PRL-02', 'PRL', 'Conducción de carretillas elevadoras', 2, date(2026, 6, 15), 'Prevención', 'Por puesto', 'Logística; Almacén', 24),
    ('PRL-03', 'PRL', 'Bloqueo y consignación de energías (LOTO)', 1, date(2026, 3, 10), 'Prevención', 'Por puesto', 'Mantenimiento; Producción', 12),
    ('PRL-04', 'PRL', 'Manipulación de productos químicos', 2, date(2026, 8, 20), 'Prevención', 'Por puesto', 'Producción', 0),
    ('IT-10', 'IT', 'Arranque y ajuste de la prensa 2', 4, date(2026, 10, 1), 'Ingeniería', 'Por puesto', 'Producción', 0),
    ('IT-11', 'IT', 'Cambio de formato en línea de corte', 2, date(2026, 7, 5), 'Ingeniería', 'Por puesto', 'Producción', 0),
    ('CAL-01', 'Calidad', 'Control de primera pieza', 1, date(2026, 5, 2), 'Calidad', 'Por puesto', 'Producción; Calidad', 0),
    ('POL-01', 'Política', 'Política de seguridad y salud', 1, date(2026, 1, 15), 'Dirección', 'Toda la plantilla', '', 0),
]
fila_fin_i = I0 + MAX_INSTR
zona_input(wi, I0 + 1, fila_fin_i, 1, 9)
for i, row in enumerate(ejemplo_instr):
    for j, v in enumerate(row, start=1):
        wi.cell(row=I0 + 1 + i, column=j, value=v)
for rr in range(I0 + 1, fila_fin_i + 1):
    wi.cell(row=rr, column=5).number_format = 'DD/MM/YYYY'
dv_tipo = DataValidation(type='list', formula1='"PRL,IT,Calidad,Política,Otra"', allow_blank=True)
dv_alc = DataValidation(type='list', formula1='"Toda la plantilla,Por puesto"', allow_blank=True)
dv_ver = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='1', allow_blank=True,
                        error='La versión es un número entero (1, 2, 3…)', showErrorMessage=True)
dv_mes = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
for dv, col in ((dv_tipo, 'B'), (dv_alc, 'G'), (dv_ver, 'D'), (dv_mes, 'I')):
    wi.add_data_validation(dv); dv.add(f'{col}{I0 + 1}:{col}{fila_fin_i}')
wi.freeze_panes = wi.cell(row=I0 + 1, column=2)
wi['D4'].comment = Comment('Súbela al publicar una revisión: quienes tenían la anterior pasan solos a Pendiente.', 'REELEVO')

# ── 3. Trabajadores ──────────────────────────────────────────────────────────
wt = wb.create_sheet('Trabajadores')
titulo(wt, 'Trabajadores', 'Una fila por persona. El puesto tiene que escribirse igual que en «Puestos afectados». Cuando alguien se va, Activo = No (no borres la fila).')
T0 = 4
cabecera(wt, T0, ['ID', 'Nombre', 'Puesto', 'Turno', 'ETT', 'Fecha de alta', 'Activo', 'Pendientes', 'Caducados', 'Caducan pronto'],
         [9, 26, 16, 10, 7, 13, 8, 11, 11, 12])
ejemplo_trab = [
    ('T001', 'Ana López', 'Producción', 'Mañana', 'No', date(2021, 3, 1), 'Sí'),
    ('T002', 'Carlos Ruiz', 'Producción', 'Tarde', 'No', date(2019, 9, 16), 'Sí'),
    ('T003', 'Marta Gil', 'Producción', 'Noche', 'No', date(2023, 1, 9), 'Sí'),
    ('T004', 'Luis Pardo', 'Mantenimiento', 'Mañana', 'No', date(2018, 5, 2), 'Sí'),
    ('T005', 'Sara Vidal', 'Calidad', 'Mañana', 'No', date(2022, 11, 7), 'Sí'),
    ('T006', 'Jorge Martín', 'Logística', 'Tarde', 'No', date(2020, 2, 3), 'Sí'),
    ('T007', 'Iván Castro', 'Almacén', 'Mañana', 'No', date(2024, 4, 22), 'Sí'),
    ('T008', 'Nerea Sanz', 'Producción', 'Noche', 'Sí', date(2026, 9, 29), 'Sí'),
    ('T009', 'Pablo Rey', 'Producción', 'Tarde', 'Sí', date(2026, 10, 6), 'Sí'),
    ('T010', 'Elena Mora', 'Producción', 'Mañana', 'No', date(2017, 6, 12), 'No'),
]
fila_fin_t = T0 + MAX_TRAB
zona_input(wt, T0 + 1, fila_fin_t, 1, 7)
for i, row in enumerate(ejemplo_trab):
    for j, v in enumerate(row, start=1):
        wt.cell(row=T0 + 1 + i, column=j, value=v)
for rr in range(T0 + 1, fila_fin_t + 1):
    wt.cell(row=rr, column=6).number_format = 'DD/MM/YYYY'
dv_si = DataValidation(type='list', formula1='"Sí,No"', allow_blank=True)
wt.add_data_validation(dv_si); dv_si.add(f'E{T0 + 1}:E{fila_fin_t}'); dv_si.add(f'G{T0 + 1}:G{fila_fin_t}')
wt.freeze_panes = wt.cell(row=T0 + 1, column=3)

# ── 4. Acuses ────────────────────────────────────────────────────────────────
wa = wb.create_sheet('Acuses')
titulo(wa, 'Registro de acuses', 'Una fila cada vez que alguien lee y acepta una versión. Solo se añaden filas; las antiguas son el histórico.')
A0 = 4
cabecera(wa, A0, ['ID trabajador', 'Código instrucción', 'Versión aceptada', 'Fecha de acuse', 'Test superado', 'Vía', 'Observaciones'],
         [13, 15, 12, 13, 12, 18, 40])
ejemplo_acuses = [
    # PRL-01 v3 (vigente) y anteriores
    ('T001', 'PRL-01', 3, date(2026, 9, 3)), ('T002', 'PRL-01', 2, date(2025, 9, 10)), ('T003', 'PRL-01', 3, date(2026, 9, 12)),
    ('T004', 'PRL-01', 3, date(2026, 9, 4)), ('T005', 'PRL-01', 3, date(2026, 9, 2)), ('T006', 'PRL-01', 2, date(2025, 10, 1)),
    ('T007', 'PRL-01', 3, date(2026, 9, 8)), ('T008', 'PRL-01', 3, date(2026, 9, 29)),
    ('T002', 'PRL-01', 1, date(2024, 9, 2)),
    # carretillas
    ('T006', 'PRL-02', 2, date(2024, 9, 10)), ('T007', 'PRL-02', 2, date(2026, 6, 20)),
    # LOTO
    ('T004', 'PRL-03', 1, date(2025, 10, 30)), ('T001', 'PRL-03', 1, date(2026, 3, 12)), ('T002', 'PRL-03', 1, date(2026, 3, 12)),
    # quimicos
    ('T001', 'PRL-04', 2, date(2026, 8, 22)), ('T002', 'PRL-04', 1, date(2025, 2, 1)), ('T003', 'PRL-04', 2, date(2026, 8, 25)),
    # IT prensa v4
    ('T001', 'IT-10', 4, date(2026, 10, 2)), ('T002', 'IT-10', 3, date(2026, 4, 14)), ('T003', 'IT-10', 3, date(2026, 4, 15)),
    # cambio de formato
    ('T001', 'IT-11', 2, date(2026, 7, 6)), ('T002', 'IT-11', 2, date(2026, 7, 7)), ('T003', 'IT-11', 2, date(2026, 7, 9)),
    ('T008', 'IT-11', 2, date(2026, 10, 1)),
    # calidad
    ('T005', 'CAL-01', 1, date(2026, 5, 3)), ('T001', 'CAL-01', 1, date(2026, 5, 4)),
    # politica
    ('T001', 'POL-01', 1, date(2026, 1, 20)), ('T002', 'POL-01', 1, date(2026, 1, 20)), ('T004', 'POL-01', 1, date(2026, 1, 21)),
    ('T005', 'POL-01', 1, date(2026, 1, 21)), ('T006', 'POL-01', 1, date(2026, 1, 22)),
]
fila_fin_a = A0 + MAX_ACUSES
zona_input(wa, A0 + 1, A0 + 400, 1, 7)  # rellenar el resto no hace falta: las formulas miran hasta MAX_ACUSES
for i, (t, c, v, d) in enumerate(ejemplo_acuses):
    vals = (t, c, v, d, 'Sí' if c.startswith('PRL') else 'N/A', 'Hoja de firmas', '')
    for j, val in enumerate(vals, start=1):
        wa.cell(row=A0 + 1 + i, column=j, value=val)
for rr in range(A0 + 1, fila_fin_a + 1):
    wa.cell(row=rr, column=4).number_format = 'DD/MM/YYYY'
dv_test = DataValidation(type='list', formula1='"Sí,No,N/A"', allow_blank=True)
dv_via = DataValidation(type='list', formula1='"Hoja de firmas,Sesión formativa,Correo,Aplicación,Otra"', allow_blank=True)
wa.add_data_validation(dv_test); dv_test.add(f'E{A0 + 1}:E{fila_fin_a}')
wa.add_data_validation(dv_via); dv_via.add(f'F{A0 + 1}:F{fila_fin_a}')
wa.freeze_panes = wa.cell(row=A0 + 1, column=1)

# ── 5. Resumen (solo la fecha de referencia aqui; los totales despues de la matriz)
wr = wb.create_sheet('Resumen', 1)
REF = "Resumen!$C$4"

# ── 6. Matriz ────────────────────────────────────────────────────────────────
wm = wb.create_sheet('Matriz')
titulo(wm, 'Matriz de estado (se calcula sola)', 'No escribas aquí. OK · Pendiente · Caducado · Caduca pronto · «—» no le afecta. Las filas son los trabajadores activos y las columnas las instrucciones.')
# Filas de cabecera de la matriz
M_COD, M_VER, M_ALC, M_PUE, M_MES = 4, 5, 6, 7, 8
M0 = 10  # primera fila de trabajadores (la 9 es su cabecera)
C0 = 4  # primera columna de instrucciones (D)
for rr, et in ((M_COD, 'Instrucción'), (M_VER, 'Versión vigente'), (M_ALC, 'Alcance'), (M_PUE, 'Puestos'), (M_MES, 'Reciclaje (meses)')):
    c = wm.cell(row=rr, column=3, value=et); c.font = font(9, True, GRIS); c.alignment = Alignment(horizontal='right')
for col, et, w in ((1, 'ID', 9), (2, 'Nombre', 24), (3, 'Puesto', 16)):
    c = wm.cell(row=M0 - 1, column=col, value=et)
    c.font = font(10, True, 'FFFFFF'); c.fill = cab_fill; c.alignment = centro
    wm.column_dimensions[get_column_letter(col)].width = w
INS = 'Instrucciones'
for j in range(MAX_INSTR):
    col = C0 + j; L = get_column_letter(col); src = I0 + 1 + j
    wm.cell(row=M_COD, column=col, value=f'=IF({INS}!$A${src}="","",{INS}!$A${src})')
    wm.cell(row=M_VER, column=col, value=f'=IF({L}${M_COD}="","",{INS}!$D${src})')
    wm.cell(row=M_ALC, column=col, value=f'=IF({L}${M_COD}="","",{INS}!$G${src})')
    wm.cell(row=M_PUE, column=col, value=f'=IF({L}${M_COD}="","",{INS}!$H${src}&"")')
    wm.cell(row=M_MES, column=col, value=f'=IF({L}${M_COD}="","",N({INS}!$I${src}))')
    for rr in (M_COD,):
        c = wm.cell(row=rr, column=col); c.font = font(10, True, 'FFFFFF'); c.fill = cab_fill; c.alignment = centro
    for rr in (M_VER, M_ALC, M_PUE, M_MES):
        c = wm.cell(row=rr, column=col); c.font = font(8, False, GRIS); c.alignment = centro; c.fill = gris_fill
    wm.cell(row=M0 - 1, column=col).fill = cab_fill
    wm.column_dimensions[L].width = 12
AC = f"Acuses!$A${A0 + 1}:$A${fila_fin_a}"
AI = f"Acuses!$B${A0 + 1}:$B${fila_fin_a}"
AV = f"Acuses!$C${A0 + 1}:$C${fila_fin_a}"
AD = f"Acuses!$D${A0 + 1}:$D${fila_fin_a}"
TR = 'Trabajadores'
for i in range(MAX_TRAB):
    rr = M0 + i; src = T0 + 1 + i
    activo = f'AND({TR}!$A${src}<>"",{TR}!$G${src}<>"No")'
    wm.cell(row=rr, column=1, value=f'=IF({activo},{TR}!$A${src},"")')
    wm.cell(row=rr, column=2, value=f'=IF($A{rr}="","",{TR}!$B${src})')
    wm.cell(row=rr, column=3, value=f'=IF($A{rr}="","",{TR}!$C${src})')
    for c in range(1, 4):
        wm.cell(row=rr, column=c).font = font(9); wm.cell(row=rr, column=c).border = borde
    for j in range(MAX_INSTR):
        col = C0 + j; L = get_column_letter(col)
        cod, ver, alc, pue, mes = (f'{L}${M_COD}', f'{L}${M_VER}', f'{L}${M_ALC}', f'{L}${M_PUE}', f'{L}${M_MES}')
        afecta = f'OR({alc}="Toda la plantilla",ISNUMBER(SEARCH(";"&SUBSTITUTE($C{rr}," ","")&";",";"&SUBSTITUTE({pue}," ","")&";")))'
        vacc = f'_xlfn.MAXIFS({AV},{AC},$A{rr},{AI},{cod})'
        facc = f'_xlfn.MAXIFS({AD},{AC},$A{rr},{AI},{cod},{AV},{ver})'
        f = (f'=IF(OR($A{rr}="",{cod}=""),"",IF(NOT({afecta}),"—",IF({vacc}<{ver},"Pendiente",'
             f'IF({mes}=0,"OK",IF(EDATE({facc},{mes})<{REF},"Caducado",IF(EDATE({facc},{mes})<{REF}+30,"Caduca pronto","OK"))))))')
        c = wm.cell(row=rr, column=col, value=f); c.font = font(9); c.alignment = centro; c.border = borde
ultima_m = M0 + MAX_TRAB - 1
rango_estado = f'{get_column_letter(C0)}{M0}:{get_column_letter(C0 + MAX_INSTR - 1)}{ultima_m}'
for txt, color, fcol in (('Pendiente', 'FDE2D8', 'B3360C'), ('Caducado', 'F8C9C9', '9C0006'), ('Caduca pronto', 'FFF0C2', '7A5A00'), ('OK', 'E2F2E2', '1E6B1E')):
    wm.conditional_formatting.add(rango_estado, CellIsRule(operator='equal', formula=[f'"{txt}"'], fill=PatternFill('solid', fgColor=color), font=Font(name=F, size=9, color=fcol, bold=txt != 'OK')))
wm.freeze_panes = wm.cell(row=M0, column=C0)
wm.sheet_view.zoomScale = 90

# ── Totales por instruccion y por trabajador ────────────────────────────────
for j in range(MAX_INSTR):
    src = I0 + 1 + j; L = get_column_letter(C0 + j)
    rng = f'Matriz!${L}${M0}:${L}${ultima_m}'
    wi.cell(row=src, column=10, value=f'=IF($A{src}="","",COUNTIF({rng},"OK")+COUNTIF({rng},"Pendiente")+COUNTIF({rng},"Caducado")+COUNTIF({rng},"Caduca pronto"))')
    wi.cell(row=src, column=11, value=f'=IF($A{src}="","",COUNTIF({rng},"Pendiente"))')
    wi.cell(row=src, column=12, value=f'=IF($A{src}="","",COUNTIF({rng},"Caducado"))')
    wi.cell(row=src, column=13, value=f'=IF(OR($A{src}="",N(J{src})=0),"",(COUNTIF({rng},"OK")+COUNTIF({rng},"Caduca pronto"))/J{src})')
    wi.cell(row=src, column=13).number_format = '0%'
    for c in range(10, 14):
        wi.cell(row=src, column=c).border = borde; wi.cell(row=src, column=c).font = font(10); wi.cell(row=src, column=c).alignment = Alignment(horizontal='center')
rng_fila = lambda rr: f'Matriz!${get_column_letter(C0)}${rr}:${get_column_letter(C0 + MAX_INSTR - 1)}${rr}'
for i in range(MAX_TRAB):
    src = T0 + 1 + i; rr = M0 + i
    wt.cell(row=src, column=8, value=f'=IF(Matriz!$A${rr}="","",COUNTIF({rng_fila(rr)},"Pendiente"))')
    wt.cell(row=src, column=9, value=f'=IF(Matriz!$A${rr}="","",COUNTIF({rng_fila(rr)},"Caducado"))')
    wt.cell(row=src, column=10, value=f'=IF(Matriz!$A${rr}="","",COUNTIF({rng_fila(rr)},"Caduca pronto"))')
    for c in range(8, 11):
        wt.cell(row=src, column=c).border = borde; wt.cell(row=src, column=c).font = font(10); wt.cell(row=src, column=c).alignment = Alignment(horizontal='center')
for ws_, col_, rng_ in ((wi, 'K', f'K{I0+1}:K{fila_fin_i}'), (wi, 'L', f'L{I0+1}:L{fila_fin_i}'), (wt, 'H', f'H{T0+1}:H{fila_fin_t}'), (wt, 'I', f'I{T0+1}:I{fila_fin_t}')):
    ws_.conditional_formatting.add(rng_, CellIsRule(operator='greaterThan', formula=['0'], fill=PatternFill('solid', fgColor='FDE2D8'), font=Font(name=F, size=10, bold=True, color='B3360C')))

# ── Resumen ──────────────────────────────────────────────────────────────────
titulo(wr, 'Resumen', 'Se calcula solo a partir de Instrucciones, Trabajadores y Acuses.')
wr.column_dimensions['A'].width = 3; wr.column_dimensions['B'].width = 46; wr.column_dimensions['C'].width = 16; wr.column_dimensions['D'].width = 70
wr['B4'] = 'Fecha de referencia'; wr['B4'].font = font(10, True)
wr['C4'] = '=TODAY()'; wr['C4'].number_format = 'DD/MM/YYYY'; wr['C4'].fill = input_fill; wr['C4'].font = font(10, True); wr['C4'].border = borde
wr['D4'] = 'Por defecto, hoy. Escribe otra fecha para ver cómo estabas ese día (p. ej., el de una auditoría).'; wr['D4'].font = font(9, False, GRIS, True); wr['D4'].alignment = Alignment(wrap_text=True, vertical='center'); wr.row_dimensions[4].height = 26
MZ = f"Matriz!{rango_estado.split(':')[0]}:{rango_estado.split(':')[1]}"
MZ = 'Matriz!$' + rango_estado.replace(':', ':$').replace(get_column_letter(C0) + str(M0), f'{get_column_letter(C0)}${M0}', 1)
c1 = f'{get_column_letter(C0)}{M0}'; c2 = f'{get_column_letter(C0 + MAX_INSTR - 1)}{ultima_m}'
MZ = f"Matriz!${get_column_letter(C0)}${M0}:${get_column_letter(C0 + MAX_INSTR - 1)}${ultima_m}"
filas = [
    ('Instrucciones registradas', f'=COUNTA(Instrucciones!$A${I0+1}:$A${fila_fin_i})', 'Todas las que tienen código en la hoja Instrucciones.'),
    ('Trabajadores activos', f'=COUNTIF(Matriz!$A${M0}:$A${ultima_m},"?*")', 'Los que tienen Activo distinto de «No».'),
    ('Personal de ETT activo', f'=COUNTIFS(Trabajadores!$E${T0+1}:$E${fila_fin_t},"Sí",Trabajadores!$G${T0+1}:$G${fila_fin_t},"<>No")', ''),
    ('Combinaciones persona-instrucción a controlar', f'=COUNTIF({MZ},"OK")+COUNTIF({MZ},"Pendiente")+COUNTIF({MZ},"Caducado")+COUNTIF({MZ},"Caduca pronto")', 'Cada una es algo que alguien tiene que poder demostrar. Crece mucho más deprisa que la plantilla.'),
    ('Al día', f'=COUNTIF({MZ},"OK")+COUNTIF({MZ},"Caduca pronto")', ''),
    ('Pendientes (no han aceptado la versión vigente)', f'=COUNTIF({MZ},"Pendiente")', 'Cada revisión de una instrucción vuelve a dejar pendientes a todos sus afectados.'),
    ('Caducados (reciclaje vencido)', f'=COUNTIF({MZ},"Caducado")', ''),
    ('Caducan en los próximos 30 días', f'=COUNTIF({MZ},"Caduca pronto")', ''),
    ('% al día', '=IF(C9=0,"",C10/C9)', 'Objetivo: 100 % en las instrucciones de seguridad.'),
]
for k, (et, fo, nota) in enumerate(filas):
    rr = 6 + k
    wr.cell(row=rr, column=2, value=et).font = font(10, k in (3, 5, 6))
    c = wr.cell(row=rr, column=3, value=fo); c.font = font(12, True, NARANJA if k in (5, 6) else TINTA); c.alignment = Alignment(horizontal='center'); c.border = borde
    n = wr.cell(row=rr, column=4, value=nota); n.font = font(9, False, GRIS, True); n.alignment = Alignment(wrap_text=True, vertical='center')
    if len(nota) > 75: wr.row_dimensions[rr].height = 26
wr['C14'].number_format = '0%'
wr['B17'] = '¿Te salen muchos pendientes cada vez que cambias una versión?'; wr['B17'].font = font(11, True, NARANJA)
wr['B18'] = ('Es normal: es aritmética, no falta de diligencia. Por qué la gestión documental se desborda con carpetas y Excel, '
             'y cómo se mantiene sola la lista de pendientes: gmvsolutions.es/blog/gestion-instrucciones-prl-e-it/')
wr['B18'].font = font(10); wr['B18'].alignment = wrap; wr.merge_cells('B18:D19'); wr.row_dimensions[18].height = 30
wr.sheet_view.showGridLines = False

for w_ in wb.worksheets:
    w_.sheet_properties.tabColor = NARANJA if w_.title in ('Cómo usar', 'Resumen') else ('FFC000' if w_.title in ('Instrucciones', 'Trabajadores', 'Acuses') else '7F7F7F')
wb.active = 0
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print('Generado', OUT)
